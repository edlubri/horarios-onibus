from pathlib import Path
import json
import re
from datetime import datetime, time
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "horarios_coletivo.xlsx"
OUT = ROOT / "public" / "data" / "linhas.json"

OUT.parent.mkdir(parents=True, exist_ok=True)


def clean(value):
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%H:%M")

    if isinstance(value, time):
        return value.strftime("%H:%M")

    return str(value).strip()


def sheet_key(name):
    s = name.strip()

    m = re.fullmatch(r"(\d{3})\s*\((Extra)\)", s, re.I)
    if m:
        return f"{m.group(1)} Extra"

    m = re.fullmatch(r"(\d{3})", s)
    if m:
        return m.group(1)

    return None


def eh_marcador_periodo(texto):
    if not texto:
        return False

    t = str(texto).strip().upper()

    return (
        "SEGUNDA-FEIRA" in t
        or "SEGUNDA FEIRA" in t
        or t in ("SÁBADO", "SABADO", "SÁBADOS", "SABADOS")
        or t in ("DOMINGO", "DOMINGOS")
        or "TERÇA-FEIRA" in t
        or "TERCA-FEIRA" in t
        or "QUARTA-FEIRA" in t
        or "QUARTA FEIRA" in t
        or "QUINTA-FEIRA" in t
        or "QUINTA FEIRA" in t
    )


def eh_horario(valor):
    if not valor:
        return False

    if valor == "/":
        return True

    return bool(re.fullmatch(r"\d{1,2}:\d{2}", valor))


def limpar_linha(row):
    return [clean(v) for v in row]


wb = openpyxl.load_workbook(
    XLSX,
    data_only=True,
    read_only=True
)

lines = []


for nome_aba in wb.sheetnames:

    codigo = sheet_key(nome_aba)

    if not codigo:
        continue

    ws = wb[nome_aba]

    rows = [
        limpar_linha(row)
        for row in ws.iter_rows(values_only=True)
    ]

    # ---------------------------------------------------------
    # TÍTULO E TRAJETO
    # ---------------------------------------------------------

    titulo = ""

    for row in rows[:5]:
        if row and row[0]:
            titulo = row[0]
            break

    trajeto = rows[2][0] if len(rows) > 2 and rows[2] else ""

    periodos = []

    # ---------------------------------------------------------
    # LOCALIZAR BLOCOS DE HORÁRIOS
    # ---------------------------------------------------------

    i = 0

    while i < len(rows):

        row = rows[i]

        if not row:
            i += 1
            continue

        primeiro = row[0]

        # A partir de EXCEÇÕES não procuramos mais tabelas.
        if str(primeiro).strip().upper().startswith("EXCE"):
            break

        if not eh_marcador_periodo(primeiro):
            i += 1
            continue

        nome_periodo = primeiro

        # Próxima linha não vazia = cabeçalho
        h = i + 1

        while h < len(rows) and not rows[h]:
            h += 1

        if h >= len(rows):
            break

        colunas = list(rows[h])

        # Remover colunas vazias no final
        while colunas and colunas[-1] == "":
            colunas.pop()

        if len(colunas) < 2:
            i += 1
            continue

        n = len(colunas)

        horarios = []

        # -----------------------------------------------------
        # LER HORÁRIOS ATÉ ENCONTRAR FIM DO BLOCO
        # -----------------------------------------------------

        j = h + 1

        while j < len(rows):

            atual = rows[j]

            # Linha vazia encerra o bloco
            if not atual:
                break

            primeiro_atual = atual[0]

            # Novo período encerra o bloco atual
            if eh_marcador_periodo(primeiro_atual):
                break

            # Exceções encerram a leitura
            if str(primeiro_atual).strip().upper().startswith("EXCE"):
                break

            valores = (atual + [""] * n)[:n]

            # Só aceitar linhas que realmente contenham
            # horários ou "/"
            if not any(eh_horario(v) for v in valores):
                break

            horarios.append(valores)

            j += 1

        if horarios:

            periodos.append({
                "nome": nome_periodo,
                "colunas": colunas,
                "horarios": horarios
            })

        # Continuar a partir do fim deste bloco
        i = j

    # ---------------------------------------------------------
    # COMPATIBILIDADE COM O SITE
    # ---------------------------------------------------------

    if periodos:

        colunas = periodos[0]["colunas"]

        todos_horarios = []

        for periodo in periodos:
            todos_horarios.extend(periodo["horarios"])

        dias = periodos[0]["nome"]

    else:

        colunas = []
        todos_horarios = []
        dias = ""

    lines.append({
        "id": codigo.lower().replace(" ", "-"),
        "codigo": codigo,
        "titulo": titulo,
        "trajeto": trajeto,
        "dias": dias,
        "colunas": colunas,
        "horarios": todos_horarios,
        "periodos": periodos,
        "aba_original": nome_aba.strip()
    })


# -------------------------------------------------------------
# ORDENAÇÃO
# -------------------------------------------------------------

def chave_ordenacao(item):

    m = re.match(r"(\d+)", item["codigo"])

    numero = int(m.group(1)) if m else 9999

    extra = 1 if "Extra" in item["codigo"] else 0

    return numero, extra


lines.sort(key=chave_ordenacao)


# -------------------------------------------------------------
# GERAR JSON
# -------------------------------------------------------------

payload = {
    "versao": 2,
    "fonte": "horarios_coletivo.xlsx",
    "total_linhas": len(lines),
    "linhas": lines
}

OUT.write_text(
    json.dumps(
        payload,
        ensure_ascii=False,
        indent=2
    ),
    encoding="utf-8"
)

print(f"Geradas {len(lines)} linhas em {OUT}")

